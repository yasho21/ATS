"""Resume text extraction from PDF and DOCX files."""
from __future__ import annotations

import io

import docx
from pypdf import PdfReader

ALLOWED_EXTENSIONS = {"pdf", "docx"}


class UnsupportedFileError(ValueError):
    pass


def get_extension(filename: str) -> str:
    return filename.rsplit(".", 1)[-1].lower() if "." in filename else ""


def extract_text(file_bytes: bytes, filename: str) -> str:
    ext = get_extension(filename)
    if ext == "pdf":
        return _extract_pdf(file_bytes)
    if ext == "docx":
        return _extract_docx(file_bytes)
    raise UnsupportedFileError(
        f"Unsupported file type '.{ext}'. Please upload a PDF or DOCX file."
    )


def _extract_pdf(file_bytes: bytes) -> str:
    reader = PdfReader(io.BytesIO(file_bytes))
    pages_text = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(pages_text)


def _extract_docx(file_bytes: bytes) -> str:
    document = docx.Document(io.BytesIO(file_bytes))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)
